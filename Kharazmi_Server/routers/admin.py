from pydantic import BaseModel
from fastapi import APIRouter, Depends, HTTPException, status, File, UploadFile, Header
from sqlalchemy.orm import Session, joinedload
from typing import List, Optional, Union
from sqlalchemy import desc, or_, text, func, case
import io
import uuid
import os
import datetime
import re
import secrets  # FIX H12: پسورد تصادفی امن (الگوی C12)

import models
from models import (
    Attendance, Course, Enrollment, Grade, InstituteShare, SessionLog, SmsLog, Student, Teacher, Transaction, User, UserSession, InstituteSettings, ActivityLog, Notification, DeviceToken, Room, Conversation, ConversationParticipant, Message, Exam, ExamQuestion, ExamAttempt, Lead
)
from schemas import (
    HistoryRequest, LoginRequest, TeacherInfo, FullTeacherProfile, StudentCreate, TeacherCreate, CourseCreate, EnrollmentCreate, GradeCreate, GradeItem, AttendanceLogRequest, AttendanceItem, AttendanceSubmitData, SmsSendRequest, ChangePasswordRequest, StudentProfileInfo, FullStudentProfile, TeacherProfileInfo, FullTeacherProfile, ClassReportInfo, ClassStudentData, ClassSessionHistory, FullClassReport, ShareConfigModel, StudentUpdate, TeacherUpdate, PersonListItem, TransactionUpdate, StudentAttendanceHistoryRequest, AdvancedSearchItem, FinanceSubmitData, PrintReceiptRequest, TransactionTestData, BulkSmsRequest, BulkSuspendRequest, InstituteSettingsModel, PricingTableUpdateModel, TeacherListItem, SmsHistoryItem
)
from dependencies import get_db, check_admin_access, check_admin_or_secretary_access, check_user_login, check_student_access, get_enrollment_tuition_and_discount, SESSION_EXPIRY_DAYS, get_current_user, hash_password, verify_password, get_next_sequence_value, normalize_mobile

# FIX: Bug 16 - share the tuition-minus-payment debt calculation across financial views.
from financial_calculations import calculate_student_debt, MAX_TEACHER_SESSION_PRICE

router = APIRouter()


@router.get("/admin/today_summary")
def get_admin_today_summary(
    db: Session = Depends(get_db),
    sub_role: str = Depends(check_user_login),
):
    """Real-time read-only summary for the admin dashboard."""
    if sub_role not in ("admin", "secretary"):
        raise HTTPException(
            status_code=403,
            detail="شما دسترسی لازم برای مشاهده وضعیت امروز را ندارید",
        )

    from today_summary import build_admin_today_summary

    return build_admin_today_summary(db=db, sub_role=sub_role)


@router.get("/dashboard/stats")
def get_dashboard_stats(db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):  # FIX (L14/Y3): آمار داشبورد — فقط کارکنان
    student_count = db.query(Student).filter(Student.is_deleted == False).count()
    class_count = db.query(Course).count()

    # FIX: Bug 12 - exclude archived Transaction rows from this active view.
    last_trans = db.query(Transaction).filter(Transaction.is_deleted == False, Transaction.is_reversed == False).order_by(desc(Transaction.id)).first()
    last_trans_data = None
    if last_trans:
        enrollment = (
            db.query(Enrollment)
            .filter(Enrollment.id == last_trans.enrollment_id)
            .first()
        )
        student_name = "ناشناس"
        if enrollment:
            st = db.query(Student).filter(Student.id == enrollment.student_id).first()
            if st:
                student_name = f"{st.first_name} {st.last_name}"
        last_trans_data = {
            "student_name": student_name,
            "amount": last_trans.amount,
            "date": last_trans.date,
        }

    last_course = db.query(Course).order_by(desc(Course.id)).first()
    last_course_data = None
    if last_course:
        teacher = db.query(Teacher).filter(Teacher.id == last_course.teacher_id).first()
        teacher_name = (
            f"{teacher.first_name} {teacher.last_name}" if teacher else "نامشخص"
        )
        last_course_data = {
            "title": last_course.title,
            "code": last_course.code,
            "teacher": teacher_name,
            "grade": last_course.grade_level,
        }

    return {
        "student_count": student_count,
        "class_count": class_count,
        "last_transaction": last_trans_data,
        "last_course": last_course_data,
    }


@router.get("/test/debt_calculation")
def test_debt_calculation(db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    """Test endpoint to verify debt calculation logic"""
    test_student = db.query(Student).first()
    if not test_student:
        return {"error": "No students found in database"}

    # Calculate debt using the new logic
    w_t = test_student.wallet_teacher if test_student.wallet_teacher is not None else 0
    w_i = (
        test_student.wallet_institute
        if test_student.wallet_institute is not None
        else 0
    )

    debt_teacher = abs(w_t) if w_t < 0 else 0
    debt_institute = abs(w_i) if w_i < 0 else 0
    # FIX: Bug 16 - this diagnostic must use the same formula as production debt reports.
    total_debt = calculate_student_debt(db, test_student)

    return {
        "student_id": test_student.id,
        "student_name": f"{test_student.first_name} {test_student.last_name}",
        "wallet_teacher": w_t,
        "wallet_institute": w_i,
        "wallet_balance": test_student.wallet_balance,
        "calculated_debt_teacher": debt_teacher,
        "calculated_debt_institute": debt_institute,
        "calculated_total_debt": total_debt,
        # FIX: Bug 16 - describe the contractual balance and the limited legacy fallback.
        "note": "Debt is discounted tuition minus recorded payments; wallets are an unpriced legacy fallback.",
    }


# ==========================================
# 6. API های حضور و غیاب
# ==========================================
class AttendanceLogRequest(BaseModel):
    course_id: int
    date: str


class AttendanceItem(BaseModel):
    student_id: int
    status: str


class AttendanceSubmitData(BaseModel):
    course_id: int
    date: str
    items: List[AttendanceItem]


@router.post("/sms/send")
def send_sms(req: SmsSendRequest, db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):  # FIX (audit-v2/#10): ارسال گروهی فقط ادمین/منشی
    count = 0
    if req.target_group == "all_students":
        count = db.query(Student).filter(Student.is_deleted == False).count()
    elif req.target_group == "all_teachers":
        count = db.query(Teacher).filter(Teacher.is_deleted == False).count()
    elif req.target_group == "debtors":
        # FIX: Bug 16 - debtor notices must include unpaid tuition even with a positive wallet.
        count = sum(calculate_student_debt(db, student) > 0 for student in db.query(Student).filter(Student.is_deleted == False).all())
    today = datetime.datetime.now().strftime("%Y/%m/%d %H:%M")
    db.add(
        models.SmsLog(
            target_group=req.target_group,
            message_text=req.message_text,
            sent_count=count,
            date=today,
        )
    )
    db.commit()
    return {"message": "ارسال شد", "count": count}


# FIX(security): ماسک کدهای OTP و سایر رشته‌های عددی داخل متن پیامک (کدها ۶رقمی‌اند؛ سال/مبلغ هم ماسک می‌شود).
_OTP_MASK_RE = re.compile(r"\d{4,}")

def _mask_otp_codes(text: Optional[str]) -> Optional[str]:
    if not text:
        return text
    return _OTP_MASK_RE.sub("***", text)


@router.get("/sms/history", response_model=List[SmsHistoryItem])
def get_sms_history(db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):  # FIX (L14/R1-align): تاریخچه برای منشی هم باز — چون ارسال (تکی+گروهی) دارد، ندیدن تاریخچه ناسازگار بود.
    logs = db.query(models.SmsLog).order_by(desc(models.SmsLog.id)).all()
    return [
        SmsHistoryItem(
            id=log.id,
            target_group=log.target_group,
            message_text=_mask_otp_codes(log.message_text),
            sent_count=log.sent_count,
            date=log.date,
        )
        for log in logs
    ]


# ==========================================
# 8. API های مدیریت معلم (تایید)
# ==========================================


@router.get("/teachers/pending", response_model=List[TeacherListItem])
def get_pending_teachers(db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):  # FIX (audit-v2/#11): لیست درانتظار (حاوی موبایل) فقط ادمین/منشی
    # FIX(security): هرگز ORM خام برنگردان — فقط فیلدهای غیرحساس (بدون password/card_number/national_code).
    teachers = db.query(Teacher).filter(Teacher.is_deleted == False).filter(Teacher.is_approved == False).all()
    return [
        TeacherListItem(
            id=t.id,
            first_name=t.first_name or "",
            last_name=t.last_name or "",
            mobile=t.mobile,
            teacher_code=t.teacher_code,
            profile_image=t.profile_image,
            is_approved=bool(t.is_approved),
            is_suspended=bool(t.is_suspended),
        )
        for t in teachers
    ]


@router.post("/teachers/approve/{teacher_id}")
def approve_teacher(teacher_id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):  # FIX H10-S2: تایید فقط ادمین/منشی (الگوی C1/H10-S)
    teacher = db.query(Teacher).filter(Teacher.id == teacher_id, Teacher.is_deleted == False).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="معلم یافت نشد")
    teacher.is_approved = True
    db.commit()
    return {"message": "معلم با موفقیت تایید شد و اکنون می‌تواند وارد شود"}


@router.delete("/teachers/reject/{teacher_id}")
def reject_teacher(teacher_id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    teacher = db.query(Teacher).filter(Teacher.id == teacher_id, Teacher.is_deleted == False).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="معلم یافت نشد")

    # سافت‌دیلیت: ردیف حفظ می‌شود (تاریخچه + جلوگیری از ثبت‌نام تکراری با همان کد ملی/موبایل)
    teacher.is_deleted = True
    db.commit()
    return {"message": "درخواست ثبت‌نام معلم رد و آرشیو شد"}


@router.post("/admin/teachers/{teacher_id}/suspend")
def suspend_teacher(teacher_id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):  # FIX H10-S: تعلیق فقط ادمین/منشی (الگوی C1)
    teacher = db.query(Teacher).filter(Teacher.id == teacher_id, Teacher.is_deleted == False).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="معلم یافت نشد")

    # Toggle suspension status
    teacher.is_suspended = not teacher.is_suspended
    db.commit()

    status_text = "تعلیق شد" if teacher.is_suspended else "فعال شد"
    return {"message": f"معلم {status_text}", "is_suspended": teacher.is_suspended}


@router.delete("/admin/teachers/{teacher_id}")
def delete_teacher(teacher_id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    teacher = db.query(Teacher).filter(Teacher.id == teacher_id, Teacher.is_deleted == False).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="معلم یافت نشد")

    # Check if teacher has any active classes
    active_classes = db.query(Course).filter(Course.teacher_id == teacher_id).count()
    if active_classes > 0:
        raise HTTPException(
            status_code=400,
            detail=f"امکان حذف معلم وجود ندارد. این معلم {active_classes} کلاس فعال دارد.",
        )

    # سافت‌دیلیت + ابطال نشست‌ها و کاربر ورود معلم (دیگر نتواند وارد شود)
    teacher.is_deleted = True
    db.query(UserSession).filter(UserSession.sub_role == "teacher", UserSession.user_id == teacher_id).delete()
    if teacher.mobile:
        db.query(User).filter(User.username == teacher.mobile).delete()
    db.commit()
    return {"message": "معلم با موفقیت حذف (آرشیو) شد"}


@router.delete("/admin/reject_class/{course_id}")
def reject_class(course_id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    c = db.query(Course).filter(Course.id == course_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="کلاس یافت نشد")

    # حذف کلاس و تصحیح حساب مبالغ تمامی دانش‌آموزان ثبت‌نام شده
    enrollments = db.query(Enrollment).filter(Enrollment.course_id == c.id).all()
    from dependencies import perform_delete_enrollment
    for en in enrollments:
        perform_delete_enrollment(en, db)
        
    # FIX: Bug 13 - archived enrollments/transactions must retain a valid course reference.
    c.is_deleted = True
    db.commit()
    return {"message": "کلاس رد و حذف شد و ترازهای مالی اصلاح گردید."}


# ==========================================
# 10. API اختصاصی پنل معلم
# ==========================================


@router.get("/admin/students/{id}/full_profile")
def get_student_full_profile(id: int, authorization: Optional[str] = Header(None), db: Session = Depends(get_db), role: str = Depends(check_user_login)):
    # FIX (L14/Y1): بازنویسی کامل تله‌ی گارد نصفه (مشابه students.py:414) — شاگرد/ولیِ غریبه دیگر رد نمی‌شوند.
    check_student_access(id, authorization, db, role)

    st = db.query(Student).filter(Student.id == id, Student.is_deleted == False).first()
    if not st:
        raise HTTPException(404, "دانش‌آموز یافت نشد")

    # لیست کلاس‌ها
    classes = []
    for en in st.enrollments:
        # FIX: Bug 13 - cancelled enrollment history is not an active class.
        if en.course and not en.is_deleted:
            classes.append(f"{en.course.title} (کد: {en.course.code})")

    # FIX(invoice): enrollmentهای فعال با شناسه‌های واقعی — کلاینت (صدور فیش) enrollment درست را
    # انتخاب/می‌فرستد؛ دیگر حدس از روی نام نمایشی کلاس لازم نیست و backend مجبور به حدس‌زدن
    # برای دانش‌آموز چندکلاسه نمی‌شود (400 «چند ثبت‌نام فعال» فقط واقعاً مبهم می‌ماند).
    enrollments_list = [
        {
            "enrollment_id": en.id,
            "course_id": en.course_id,
            "title": en.course.title or "",
            "code": en.course.code or "",
            "branch_id": en.branch_id,
        }
        for en in st.enrollments
        if not en.is_deleted and en.course is not None
    ]

    # لیست تراکنش‌ها
    trans = (
        # FIX: Bug 12 - exclude archived Transaction rows from this active view.
        db.query(Transaction).filter(Transaction.is_deleted == False, Transaction.is_reversed == False)
        .filter(Transaction.student_id == id)
        .order_by(desc(Transaction.id))
        .limit(5)
        .all()
    )
    trans_list = [f"{t.date}: {t.amount:,} ({t.description})" for t in trans]

    # محاسبه دقیق بدهی‌ها از دیتابیس
    w_t = st.wallet_teacher if st.wallet_teacher is not None else 0
    w_i = st.wallet_institute if st.wallet_institute is not None else 0

    # اگر منفی باشد یعنی بدهکار است (برای نمایش مثبتش میکنیم)
    debt_t = abs(w_t) if w_t < 0 else 0
    debt_i = abs(w_i) if w_i < 0 else 0
    # FIX: Bug 16 - expose the tuition-based total while retaining separate wallet balances.
    total_d = calculate_student_debt(db, st)
    wallet_total = w_t + w_i

    # محاسبات تفکیک شده مالی به ازای مربیان و آموزشگاه
    teachers_financial = []
    for en in st.enrollments:
        # FIX: Bug 13 - do not include archived enrollments in financial class balances.
        if en.is_deleted or not en.course:
            continue
        course = en.course
        teacher = db.query(Teacher).filter(Teacher.id == course.teacher_id).first()
        teacher_name = f"{teacher.first_name} {teacher.last_name}" if teacher else "بدون معلم"
        
        # 1. مجموع کل پرداختی مربی برای این کلاس - فیلتر استاندارد is_deleted و is_reversed
        paid_teacher = db.query(func.sum(Transaction.amount)).filter(
            Transaction.student_id == id,
            Transaction.course_id == course.id,
            Transaction.target_wallet == "teacher",
            Transaction.amount > 0,
            Transaction.is_deleted == False,
            Transaction.is_reversed == False
        ).scalar() or 0
        
        # 2. مجموع کل هزینه برگزاری جلسات برای این مربی - فیلتر استاندارد
        billed_teacher = db.query(func.sum(Transaction.share_teacher)).filter(
            Transaction.student_id == id,
            Transaction.course_id == course.id,
            Transaction.type == "session_charge",
            Transaction.is_deleted == False,
            Transaction.is_reversed == False
        ).scalar() or 0
        
        # 3. مجموع کل پرداختی آموزشگاه برای این کلاس - فیلتر استاندارد
        paid_inst = db.query(func.sum(Transaction.amount)).filter(
            Transaction.student_id == id,
            Transaction.course_id == course.id,
            Transaction.target_wallet == "institute",
            Transaction.amount > 0,
            Transaction.is_deleted == False,
            Transaction.is_reversed == False
        ).scalar() or 0
        
        # 4. مجموع کل هزینه برگزاری جلسات برای این آموزشگاه - فیلتر استاندارد مثل calculate_institute_session_revenue
        billed_inst = db.query(func.sum(Transaction.share_institute)).filter(
            Transaction.student_id == id,
            Transaction.course_id == course.id,
            Transaction.type == "session_charge",
            Transaction.is_deleted == False,
            Transaction.is_reversed == False
        ).scalar() or 0
        
        debt_teacher_course = max(0, billed_teacher - paid_teacher)
        debt_inst_course = max(0, billed_inst - paid_inst)
        
        teachers_financial.append({
            "course_title": course.title,
            "teacher_name": teacher_name,
            "paid_teacher": int(paid_teacher),
            "debt_teacher": int(debt_teacher_course),
            "paid_institute": int(paid_inst),
            "debt_institute": int(debt_inst_course)
        })

    # مجموع کل پرداخت دانش‌آموز به آموزشگاه (همه کلاس‌ها) - فیلتر استاندارد
    total_paid_institute_overall = db.query(func.sum(Transaction.amount)).filter(
        Transaction.student_id == id,
        Transaction.target_wallet == "institute",
        Transaction.amount > 0,
        Transaction.is_deleted == False,
        Transaction.is_reversed == False
    ).scalar() or 0

    return {
        "info": {
            "name": f"{st.first_name} {st.last_name}",
            "national_code": st.national_code,
            "student_mobile": st.student_mobile,
            "parent_mobile": st.parent_mobile,
            "address": st.address if st.address else "---",
            "is_suspended": st.is_suspended if st.is_suspended else False,
            "profile_image": st.profile_image,
            "student_code": st.student_code,
        },
        "classes": classes,
        # FIX(invoice): شناسه‌های واقعی enrollmentهای فعال (سازگار با عقب: کلاینت‌های قدیمی این کلید را نمی‌خوانند)
        "enrollments": enrollments_list,
        "transactions": trans_list,
        "total_debt": total_d,
        "debt_teacher": debt_t,  # ✅ مقدار واقعی از دیتابیس
        "debt_institute": debt_i,  # ✅ مقدار واقعی از دیتابیس
        "wallet_teacher": w_t,
        "wallet_institute": w_i,
        "wallet_total": wallet_total,
        "total_paid_institute": int(total_paid_institute_overall),
        "teachers_financial": teachers_financial,
    }


# ==========================================
# 13. API پروفایل کامل معلم (Teacher Profile)
# ==========================================
class TeacherProfileInfo(BaseModel):
    name: str
    mobile: str
    national_code: str
    status: str  # فعال/غیرفعال
    profile_image: Optional[str] = None


class FullTeacherProfile(BaseModel):
    info: TeacherProfileInfo
    total_students: int
    total_revenue: int  # درآمدزایی کل
    active_classes_count: int
    classes: List[str]  # لیست نام کلاس‌ها


@router.get("/config/share")
def get_share_config(db: Session = Depends(get_db), _: str = Depends(check_admin_access)):  # FIX (L14/Y3): پیکربندی سهم — فقط ادمین (متقارن با update؛ اپ هم از منشی پنهان است)
    # تلاش برای گرفتن تنظیمات
    config = db.query(InstituteShare).first()
    if not config:
        # اگر نبود، یکی با مقادیر صفر میسازیم
        config = InstituteShare()
        db.add(config)
        db.commit()
        db.refresh(config)
    return config


@router.post("/config/share/update")
def update_share_config(data: ShareConfigModel, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    config = db.query(InstituteShare).first()
    if not config:
        config = InstituteShare()
        db.add(config)

    # آپدیت مقادیر
    config.count_1 = data.count_1
    config.count_2 = data.count_2
    config.count_3 = data.count_3
    config.count_4 = data.count_4
    config.count_5 = data.count_5
    config.count_6 = data.count_6
    config.count_7 = data.count_7
    config.count_8 = data.count_8
    config.count_9 = data.count_9
    config.count_10 = data.count_10
    config.count_11 = data.count_11
    config.count_12 = data.count_12
    config.count_13 = data.count_13
    config.count_14 = data.count_14
    config.count_15 = data.count_15

    db.commit()
    return {"message": "تنظیمات سهم آموزشگاه ذخیره شد"}


@router.get("/admin/pending_classes")
def get_pending_classes(db: Session = Depends(get_db), authorization: Optional[str] = Header(None), sub_role: str = Depends(check_user_login)):
    # FIX (L14/Y3): کلاس‌های درانتظار — کارکنان کامل؛ معلم فقط پیشنهادهای خودش (پیگیری ثبت از ClassSetup)؛ شاگرد/ولی 403.
    if sub_role not in ("admin", "secretary", "teacher"):
        raise HTTPException(status_code=403, detail="شما دسترسی لازم برای مشاهده کلاس‌های درانتظار را ندارید")
    _own_teacher_id = None
    if sub_role == "teacher":
        from routers.reports import get_logged_in_teacher  # lazy (جلوگیری از import چرخه‌ای)
        _me = get_logged_in_teacher(db, authorization)
        if not _me:
            raise HTTPException(status_code=403, detail="شما دسترسی لازم برای مشاهده کلاس‌های درانتظار را ندارید")
        _own_teacher_id = _me.id
    # 1. دریافت کلاس‌های تایید نشده از دیتابیس
    courses = db.query(Course).filter(Course.is_admin_approved == False)
    if _own_teacher_id is not None:
        courses = courses.filter(Course.teacher_id == _own_teacher_id)
    courses = courses.all()

    result = []
    for c in courses:
        # 2. پیدا کردن نام معلم برای نمایش
        teacher = db.query(Teacher).filter(Teacher.id == c.teacher_id).first()
        t_name = f"{teacher.first_name} {teacher.last_name}" if teacher else "نامشخص"

        # 3. ساخت دیکشنری دقیقاً با کلیدهایی که اندروید منتظر آن است
        # نکته کلیدی: نام سمت چپ (داخل گیومه) باید با متغیرهای AppModels.kt یکی باشد
        # FIX (audit-v2/blind-approve-2b): مبنای واقعی سهم آموزشگاه — H5 از InstituteShare؛
        # کلاسِ درانتظار شاگرد ندارد پس count_1 (مبنای یک‌شاگرد). قبلاً هاردکد ۰ بود و ادمین
        # نابینا تأیید می‌کرد. prepay یعنی واقعیِ صفر؛ نبود ردیف تعرفه هم صفر (submit جدا 500 می‌دهد).
        _base_inst_share = 0
        if not c.rule_prepay_institute:
            _share_row = db.query(models.InstituteShare).first()
            if _share_row is not None:
                _base_inst_share = getattr(_share_row, "count_1", 0) or 0
        result.append(
            {
                "id": c.id,
                "title": c.title,
                "teacher_name": t_name,  # اندروید: teacher_name
                "teacher_price": c.teacher_session_price,  # اندروید: teacher_price (در دیتابیس teacher_session_price است)
                "days": c.days_of_week,  # اندروید: days (در دیتابیس days_of_week است)
                "time": c.class_time,  # اندروید: time (در دیتابیس class_time است)
                "base_institute_share": _base_inst_share,  # FIX (audit-v2/blind-approve-2b): مبنای واقعی H5؛ صفر یعنی پیش‌پرداخت/نبود تعرفه
            }
        )

    return result


@router.post("/admin/approve_class/{course_id}")
def approve_class(course_id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):  # FIX H10-S2: تایید فقط ادمین/منشی (الگوی C1/H10-S)
    c = db.query(Course).filter(Course.id == course_id).first()
    if not c:  # FIX H10-S2: قبلاً None-check نداشت (500) — هم‌سبک reject_class همسایه
        raise HTTPException(status_code=404, detail="کلاس یافت نشد")
    c.is_admin_approved = True
    db.commit()
    # FIX (audit-v2/blind-approve-2c): هشدار غیرمسدودکننده برای نرخ بالای سقف (اشتباه تایپی
    # موقع ثبت). مسدود نکردیم: کلاس VIP مشروع هم می‌تواند بالای سقف باشد و 400 فلو تأیید اپ را
    # می‌شکست؛ ادمین با دیدن مبلغ واقعی (قدم 2b) + این هشدار، آگاهانه تأیید می‌کند.
    _price_warning = None
    if (c.teacher_session_price or 0) > MAX_TEACHER_SESSION_PRICE:
        _price_warning = f"⚠️ نرخ هر جلسه ({c.teacher_session_price:,} تومان برای هر شاگرد) بالاتر از سقف معمول ({MAX_TEACHER_SESSION_PRICE:,} تومان) است؛ اگر اشتباه تایپی است قبل از برگزاری جلسه اصلاح کنید"
    return {"message": "تایید شد", "price_warning": _price_warning}


@router.get("/admin/parent_contacts")
def get_parent_contacts(
    class_id: Optional[int] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db),
    _: str = Depends(check_admin_or_secretary_access)  # FIX (audit-v2/contacts-PII): دفترچه‌ی تماس والدین گزارش مدیریتی است — فقط ادمین/منشی
):
    q = db.query(Student).outerjoin(Enrollment).outerjoin(Course).options(
        joinedload(Student.enrollments).joinedload(Enrollment.course)
    )
    q = q.filter(Student.is_deleted == False)
    
    if class_id:
        q = q.filter(Enrollment.course_id == class_id)
        
    if search:
        search_fmt = f"%{search}%"
        q = q.filter(
            or_(
                Student.first_name.ilike(search_fmt),
                Student.last_name.ilike(search_fmt),
                Student.national_code.ilike(search_fmt)
            )
        )
        
    students = q.all()
    
    result = []
    for st in students:
        # Get student classes
        classes = []
        for en in st.enrollments:
            if en.course:
                classes.append(en.course.title)
                
        result.append({
            "student_id": st.id,
            "student_name": f"{st.first_name} {st.last_name}",
            "student_mobile": st.student_mobile or "",
            "parent_mobile": st.parent_mobile or "",
            "classes": classes
        })
    return result


# ==========================================
# Excel Export Endpoints
# ==========================================
import io
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


@router.get("/admin/transactions/list/excel")
def get_transactions_excel(search: Optional[str] = None, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    # FIX: Bug 12 - exclude archived Transaction rows from this active view.
    q = db.query(Transaction).filter(Transaction.is_deleted == False, Transaction.is_reversed == False).outerjoin(Student).order_by(desc(Transaction.id))

    if search:
        search_fmt = f"%{search}%"
        q = q.filter(
            or_(
                Student.first_name.ilike(search_fmt),
                Student.last_name.ilike(search_fmt),
                Student.national_code.ilike(search_fmt),
                Transaction.description.ilike(search_fmt),
            )
        )

    trans_list = q.all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش تراکنش‌های مالی"
    ws.views.sheetView[0].rightToLeft = True
    
    headers = ["شناسه", "نام دانش‌آموز", "کلاس", "مبلغ (تومان)", "روش پرداخت", "کد پیگیری", "تاریخ", "دریافت‌کننده", "کیف پول هدف", "سهم معلم", "سهم آموزشگاه", "توضیحات"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="00695C", end_color="00695C", fill_type="solid")
    header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    row_num = 2
    for t in trans_list:
        course_name = t.course.title if t.course else "---"
        student_name = f"{t.student.first_name} {t.student.last_name}" if t.student else "ناشناس"
        
        target_w = "معلم" if t.target_wallet == "teacher" else "آموزشگاه" if t.target_wallet == "institute" else "---"
        
        row_data = [
            t.id,
            student_name,
            course_name,
            t.amount,
            t.payment_method or "---",
            t.tracking_code or "---",
            t.date or "---",
            t.receiver or "---",
            target_w,
            t.share_teacher or 0,
            t.share_institute or 0,
            t.description or ""
        ]
        ws.append(row_data)
        
        ws.cell(row=row_num, column=4).number_format = '#,##0'
        ws.cell(row=row_num, column=10).number_format = '#,##0'
        ws.cell(row=row_num, column=11).number_format = '#,##0'
        
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Tahoma", size=10)
            
        row_num += 1
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=transactions_report.xlsx"}
    )


from sqlalchemy import or_


# مدل پاسخ برای لیست
class PersonListItem(BaseModel):
    id: int
    name: str
    national_code: str
    mobile: str
    role: str  # 'student' or 'teacher'
    is_suspended: bool = False


# 1. جستجوی دانش‌آموزان (همه یا با فیلتر)


@router.get("/admin/transactions/list")
def get_all_transactions(search: Optional[str] = None, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    # FIX: Bug 12 - exclude archived Transaction rows from this active view.
    q = db.query(Transaction).filter(Transaction.is_deleted == False, Transaction.is_reversed == False).join(Student).order_by(desc(Transaction.id))

    if search:
        search_fmt = f"%{search}%"
        q = q.filter(
            or_(
                Student.first_name.ilike(search_fmt),
                Student.last_name.ilike(search_fmt),
                Student.national_code.ilike(search_fmt),
                Transaction.description.ilike(search_fmt),
            )
        )

    trans_list = q.all()

    # مپ کردن دستی برای ارسال نام کلاس و دانش‌آموز
    result = []
    for t in trans_list:
        course_name = t.course.title if t.course else "---"
        student_name = f"{t.student.first_name} {t.student.last_name}"
        result.append(
            {
                "id": t.id,
                "student_id": t.student_id,
                "student_name": student_name,
                "course_id": t.course_id,
                "course_name": course_name,
                "amount": t.amount,
                "date": t.date,
                "description": t.description,
                "type": t.type,
            }
        )
    return result


# 2. حذف تراکنش (بسیار مهم: باید حساب کلاس هم اصلاح شود)


@router.delete("/admin/transactions/{id}")
def delete_transaction(id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    # FIX: Bug 12 - exclude archived Transaction rows from this active view.
    trans = db.query(Transaction).filter(Transaction.is_deleted == False, Transaction.is_reversed == False).filter(Transaction.id == id).first()
    if not trans:
        raise HTTPException(status_code=404, detail="تراکنش یافت نشد")

    # فریز مبلغ در لحظه‌ی اعتبارسنجی (سطر در انتهای تابع hard-delete می‌شود).
    t_amount = trans.amount

    # FIX B3 (الگوی H8-P4): گیت اتمیک گذرا — is_deleted مشروط فقط برای mutual exclusion؛ چون این تابع
    # hard-delete می‌کند (تصمیم محصولی H11، خارج از این ممیزی)، این فلیپ در state نهایی نامرئی است
    # (همان تراکنش سطر را با db.delete پاک می‌کند). rowcount≠۱ ← همان 404 حذف تکراری ترتیبی.
    claimed_rows = (
        db.query(Transaction)
        .filter(Transaction.id == id, Transaction.is_deleted == False, Transaction.is_reversed == False)
        .update({Transaction.is_deleted: True}, synchronize_session=False)
    )
    if claimed_rows != 1:
        db.rollback()
        raise HTTPException(status_code=404, detail="تراکنش یافت نشد")

    # اگر تراکنش مربوط به شهریه کلاس بوده، باید از پرداختی دانش‌آموز کم کنیم
    # 🔗 لینک دقیق ثبت‌نام مقدم است (منبع حقیقت واحد)؛ وگرنه رفتار قدیمی حفظ می‌شود
    if trans.enrollment_id and trans.type in ("deposit", "enrollment_payment", "tuition") and (trans.amount or 0) > 0:
        exact_enrollment = db.query(Enrollment).filter(Enrollment.id == trans.enrollment_id).first()
        if exact_enrollment is not None and not exact_enrollment.is_deleted:
            # FIX B3: کسر اتمیک با کف صفر (همان max(0,...) قبلی، به‌صورت CASE اتمیک) — شاخه‌ی لینک دقیق.
            db.query(Enrollment).filter(
                Enrollment.id == exact_enrollment.id, Enrollment.is_deleted == False
            ).update(
                {
                    Enrollment.total_paid: case(
                        (func.coalesce(Enrollment.total_paid, 0) - t_amount < 0, 0),
                        else_=func.coalesce(Enrollment.total_paid, 0) - t_amount,
                    )
                },
                synchronize_session=False,
            )
    elif trans.course_id and trans.type == "tuition":
        enrollment = (
            db.query(Enrollment)
            .filter(
                Enrollment.student_id == trans.student_id,
                Enrollment.course_id == trans.course_id,
            )
            .first()
        )

        if enrollment:
            # کسر مبلغ حذف شده از کل پرداختی
            # FIX B3: کسر اتمیک با کف صفر (همان floor دستی قبلی، به‌صورت CASE اتمیک) — شاخه‌ی legacy.
            # عین رفتار قبلی بدون شرط is_deleted (خواندن قدیمی هم فیلتر deleted نداشت — آینه‌ی مسیر B در update_transaction).
            db.query(Enrollment).filter(Enrollment.id == enrollment.id).update(
                {
                    Enrollment.total_paid: case(
                        (func.coalesce(Enrollment.total_paid, 0) - t_amount < 0, 0),
                        else_=func.coalesce(Enrollment.total_paid, 0) - t_amount,
                    )
                },
                synchronize_session=False,
            )

    # بروزرسانی کیف پول دانش‌آموز بر اساس نوع تراکنش
    if not trans.student_id:
        db.delete(trans)
        db.commit()
        return {"message": "تراکنش حذف شد (تراکنش بدون دانش‌آموز)"}

    student = db.query(Student).filter(Student.id == trans.student_id).first()
    if not student:
        db.delete(trans)
        db.commit()
        return {"message": "تراکنش حذف شد (دانش‌آموز یافت نشد)"}

    # FIX B3: بدون خواندن-جمع‌زدن در پایتون — فقط دلتای signed هر کیف‌پول محاسبه می‌شود (در انتها یک UPDATE اتمیک).
    # (متغیر wallet_balance قدیمی عملاً دور ریخته می‌شد — write-back فقط مؤلفه‌ها + sync بود؛ آینه‌ی update_transaction.)
    d_teacher = 0
    d_institute = 0

    if trans.type == "deposit":
        # برای تراکنش‌های واریز (مثبت)، مبلغ را از کیف پول کم می‌کنیم (برگشت واریز)
        if trans.target_wallet == "teacher":
            d_teacher = -t_amount
        elif trans.target_wallet == "institute":
            d_institute = -t_amount
        elif trans.target_wallet == "both":
            # FIX F-D1: رسیدهای legacy با target_wallet="both" (امروز submit_payment دو رسید جدا
            # می‌سازد) هنگام حذف هیچ اثری از کیف برنمی‌گرداندند ⇒ واگرایی دفتر/کیف — دقیقاً همان
            # شکافی که مسیر refund برای همین ردیف‌ها پوشش می‌دهد.
            # قاعده‌ی کانونیکال (آینه‌ی refund_transaction، بدون تغییر اعداد):
            #   جمع سهم‌های ذخیره‌شده == مبلغ ⇒ همان سهم‌ها کسر می‌شوند،
            #   وگرنه fallback مستند پروژه: نصف-نصف با باقیمانده به آموزشگاه.
            share_teacher = int(trans.share_teacher or 0)
            share_institute = int(trans.share_institute or 0)
            if share_teacher + share_institute != t_amount:
                share_teacher = t_amount // 2
                share_institute = t_amount - share_teacher
            d_teacher = -share_teacher
            d_institute = -share_institute
        # (هدف نامشخص (None): مثل رفتار قبلی هیچ تغییری اعمال نمی‌شود — wallet_balance قدیمی دور ریخته می‌شد.)

    elif trans.type == "session_charge":
        # برای تراکنش‌های هزینه جلسه (منفی)، مبلغ را به کیف پول برمی‌گردانیم
        # session_charge تراکنش‌ها مقدار منفی دارند (کسری از کیف پول)
        # برای برگرداندن باید مقدار مثبت اضافه کنیم
        share_teacher = trans.share_teacher if trans.share_teacher is not None else 0
        share_institute = (
            trans.share_institute if trans.share_institute is not None else 0
        )

        if share_teacher != 0:
            d_teacher = abs(share_teacher)
        if share_institute != 0:
            d_institute = abs(share_institute)

    elif trans.type == "tuition" or trans.type == "enrollment_payment":
        # (no-op عمدی — رفتار قبلی هم فقط wallet_balance محلی را عوض می‌کرد که دور ریخته می‌شد.)
        pass

    else:
        # (no-op عمدی — همان دلیل بالا.)
        pass

    # FIX B3 (الگوی H8-P4): اعمال اتمیک دلتاها در SQL — COALESCE + دلتای signed (منفی=کسر، مثبت=برگرداندن).
    if d_teacher or d_institute:
        wallet_rows = (
            db.query(Student)
            .filter(Student.id == trans.student_id)
            .update(
                {
                    Student.wallet_teacher: func.coalesce(Student.wallet_teacher, 0) + d_teacher,
                    Student.wallet_institute: func.coalesce(Student.wallet_institute, 0) + d_institute,
                },
                synchronize_session=False,
            )
        )
        if wallet_rows != 1:
            db.rollback()
            raise HTTPException(status_code=404, detail="دانش‌آموز یافت نشد")

    # ترتیب سینک (مثل بچ ۱/۲): اول UPDATE خام، بعد refresh برای مقادیر تازه، بعد سینک بالانس کل.
    if d_teacher or d_institute:
        db.refresh(student)
        # FIX: Bug 18 - derive the total only through the shared wallet helper.
        student.sync_wallet_balance()

    # hard-delete سر جایش (بعد از همه‌ی UPDATEهای مشروط موفق) — تصمیم H11، خارج از این ممیزی.
    db.delete(trans)
    db.commit()
    return {"message": "تراکنش حذف شد و حساب‌ها اصلاح گردید."}


# 3. ویرایش تراکنش


def _split_amount_integer(amount: int, share_a: int, share_b: int):
    """FIX F-W1a/F-W2: تقسیمِ قطعیِ یک مبلغ بین دو سهم، با جمعِ **دقیقاً** برابر amount.

    - خروجی همیشه عدد صحیح (تومان) است؛ هیچ مقدار اعشاری (float) تولید نمی‌شود.
    - قاعده‌ی قطعیِ باقیمانده: باقیمانده (**در قدر مطلق**، هم برای مبلغ مثبت و هم منفی) همیشه
      به سهم دوم (آموزشگاه) می‌رسد — همان قرارداد موجودِ پروژه در `submit_payment`
      (Bug 10: `amt_teacher = amount // 2` و `amt_institute = amount - amt_teacher`) و در مسیر
      refund رسیدهای قدیمی `both`.
    - اگر هر دو سهم صفر/غایب باشند، تقسیم نصف-نصف (۱:۱) با همان قاعده‌ی باقیمانده است —
      دقیقاً fallback مستندشده‌ی مسیر refund برای رسیدهای legacy.
    """
    share_a = int(share_a or 0)
    share_b = int(share_b or 0)
    total = abs(share_a) + abs(share_b)
    if total <= 0:
        share_a, share_b, total = 1, 1, 2
    amount = int(amount)
    sign = -1 if amount < 0 else 1
    part_a = sign * (abs(amount) * abs(share_a) // total)   # قطعی و مستقل از علامت/ترتیب فراخوانی
    return part_a, amount - part_a


@router.put("/admin/transactions/{id}")
def update_transaction(id: int, data: TransactionUpdate, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    # FIX: Bug 12 - exclude archived Transaction rows from this active view.
    trans = db.query(Transaction).filter(Transaction.is_deleted == False, Transaction.is_reversed == False).filter(Transaction.id == id).first()
    if not trans:
        raise HTTPException(status_code=404, detail="تراکنش یافت نشد")

    # FIX: Bug 22 - payment edits stay positive; signed charges/reversals keep their existing semantics.
    if trans.type in ("deposit", "enrollment_payment", "tuition") and data.amount <= 0:
        raise HTTPException(status_code=400, detail="مبلغ پرداخت باید بیشتر از صفر باشد")

    old_amount = trans.amount
    new_amount = data.amount

    # آپدیت فیلدها
    trans.amount = new_amount
    trans.description = data.description
    trans.date = data.date

    # اصلاح حساب کلاس (اگر شهریه بوده)
    # 🔗 لینک دقیق ثبت‌نام مقدم است (منبع حقیقت واحد)؛ وگرنه رفتار قدیمی حفظ می‌شود
    if trans.enrollment_id and trans.type in ("deposit", "enrollment_payment", "tuition") and (trans.amount or 0) > 0 and (new_amount or 0) > 0:
        exact_enrollment = db.query(Enrollment).filter(Enrollment.id == trans.enrollment_id).first()
        if exact_enrollment is not None and not exact_enrollment.is_deleted:
            # اختلاف مبلغ جدید و قدیم را حساب میکنیم
            diff = new_amount - old_amount
            # FIX B2 (الگوی H8-P4): افزایش/کاهش اتمیک با کف صفر (همان max(0,...) قبلی، به‌صورت CASE اتمیک).
            # diff از قبل (مبلغ جدید − مبلغ فریزشده‌ی قبل‌از‌تغییر) محاسبه شده؛ trans.amount اینجا دیگر new است.
            db.query(Enrollment).filter(
                Enrollment.id == exact_enrollment.id, Enrollment.is_deleted == False
            ).update(
                {
                    Enrollment.total_paid: case(
                        (func.coalesce(Enrollment.total_paid, 0) + diff < 0, 0),
                        else_=func.coalesce(Enrollment.total_paid, 0) + diff,
                    )
                },
                synchronize_session=False,
            )
    elif trans.course_id and trans.type == "tuition":
        enrollment = (
            db.query(Enrollment)
            .filter(
                Enrollment.student_id == trans.student_id,
                Enrollment.course_id == trans.course_id,
            )
            .first()
        )

        if enrollment:
            # اختلاف مبلغ جدید و قدیم را حساب میکنیم
            diff = new_amount - old_amount
            # FIX B2 (الگوی H8-P4): افزایش/کاهش اتمیک — عین رفتار قبلی (بدون کف صفر و بدون شرط is_deleted،
            # چون خواندن مسیر قدیمی هم هیچ‌کدام را نداشت؛ تغییر منطق محصولی در اسکوپ بچ ۲ نیست).
            db.query(Enrollment).filter(Enrollment.id == enrollment.id).update(
                {Enrollment.total_paid: func.coalesce(Enrollment.total_paid, 0) + diff},
                synchronize_session=False,
            )

    # بروزرسانی کیف پول دانش‌آموز بر اساس نوع تراکنش
    if not trans.student_id:
        db.commit()
        return {"message": "تراکنش ویرایش شد (تراکنش بدون دانش‌آموز)"}

    # FIX B2 (الگوی H8-P4): بدون with_for_update — این خواندن فقط وجودسنجی است؛ اصلاح کیف‌پول با UPDATE اتمیک پایین است.
    student = db.query(Student).filter(Student.id == trans.student_id).first()
    if not student:
        db.commit()
        return {"message": "تراکنش ویرایش شد (دانش‌آموز یافت نشد)"}

    # FIX B2: بدون خواندن-جمع‌زدن در پایتون — فقط دلتای هر کیف‌پول محاسبه می‌شود (در انتها یک UPDATE اتمیک).
    # (متغیر wallet_balance قدیمی عملاً دور ریخته می‌شد — write-back فقط مؤلفه‌ها + sync بود؛ پس حذفش بی‌اثر است.)
    d_teacher = 0
    d_institute = 0

    # diff از روی مبلغ فریزشده‌ی قبل‌از‌تغییر (old_amount) محاسبه می‌شود، نه trans.amount (که دیگر new است).
    diff = new_amount - old_amount

    if trans.type == "deposit":
        # برای تراکنش‌های واریز (مثبت)، اختلاف مبلغ را به کیف پول اضافه/کم می‌کنیم
        if trans.target_wallet == "teacher":
            d_teacher = diff
        elif trans.target_wallet == "institute":
            d_institute = diff
        elif trans.target_wallet == "both":
            # FIX F-W2: رسیدهای legacy با target_wallet="both" (امروز submit_payment دو رسید جدا
            # می‌سازد) پیش‌تر هیچ دلتایی نمی‌گرفتند ⇒ دفتر و کیف‌پول از هم واگرا می‌شدند، در حالی که
            # دو مسیر دیگر (refund و delete) همین رسیدها را می‌شناسند.
            # قاعده: delta = new_amount - old_amount، تقسیم با **همان نسبت سهم قبلی**؛
            # اگر سهم قبلی موجود/معتبر نباشد، fallback نصف-نصف (باقیمانده به آموزشگاه) — همان قاعده‌ی refund.
            prior_teacher = int(trans.share_teacher or 0)
            prior_institute = int(trans.share_institute or 0)
            if abs(prior_teacher) + abs(prior_institute) > 0:
                # نسبت قبلی موجود است ⇒ سهم‌های جدید با **همان نسبت** تقسیم می‌شوند و دلتای هر کیف
                # دقیقاً «سهم جدید − سهم قبلی» است. برای ردیف سازگار (جمع سهم قبلی == مبلغ قبلی)
                # جمع دلتاها == delta = new_amount − old_amount؛ و کیف‌ها همیشه با سهم‌های ذخیره‌شده
                # هم‌راستا می‌مانند — همان چیزی که مسیر refund (که همین ستون‌ها را برمی‌گرداند) لازم دارد.
                new_teacher, new_institute = _split_amount_integer(
                    new_amount, prior_teacher, prior_institute
                )
                d_teacher = new_teacher - prior_teacher
                d_institute = new_institute - prior_institute
                trans.share_teacher, trans.share_institute = new_teacher, new_institute
            else:
                # fallback مستند (سهم قبلی موجود نیست — ردیف‌های legacy بدون سهم): delta طبق
                # فرمول new − old نصف-نصف تقسیم می‌شود (باقیمانده به آموزشگاه) و سهم‌های ردیف هم
                # با همان نسبت ۵۰/۵۰ با مبلغ جدید هم‌راستا می‌شوند.
                d_teacher, d_institute = _split_amount_integer(diff, 1, 1)
                trans.share_teacher, trans.share_institute = _split_amount_integer(new_amount, 1, 1)
        # (هدف نامشخص (None): مثل رفتار قبلی هیچ تغییری در کیف‌پول اعمال نمی‌شود — wallet_balance قدیمی دور ریخته می‌شد.)

    elif trans.type == "session_charge":
        # FIX F-W1a/F-W1b: در session_charge قرارداد پروژه این است: amount منفی و
        # share_teacher/share_institute مثبت‌اند و `amount == -(share_teacher + share_institute)`
        # (سازنده: routers/attendance.py). ویرایش قبلی دلتا را با نسبتِ **اعشاری** حساب می‌کرد
        # (`diff * teacher_ratio`) و روی ستون BigInteger می‌نوشت ⇒ پول اعشاری؛ و share_* را هم
        # با مبلغ جدید همراستا نمی‌کرد ⇒ گزارش سهم‌ها (financial_calculations) کهنه می‌ماند.
        prior_teacher = int(trans.share_teacher or 0)
        prior_institute = int(trans.share_institute or 0)
        prior_total = abs(prior_teacher) + abs(prior_institute)

        if prior_total > 0:
            # سهم‌های جدید = تقسیم صحیحِ |مبلغ جدید| با همان نسبت قبلی (باقیمانده به آموزشگاه).
            new_teacher, new_institute = _split_amount_integer(
                abs(int(new_amount)), prior_teacher, prior_institute
            )
            # دلتای هر کیف = سهم قبلی − سهم جدید ⇒ جمع دلتاها دقیقاً برابر تغییر مبلغ است
            # و بعد از ویرایش، کسرِ قابل‌انتساب به این ردیف دقیقاً همان سهم ذخیره‌شده می‌شود.
            d_teacher = prior_teacher - new_teacher
            d_institute = prior_institute - new_institute
            trans.share_teacher = new_teacher
            trans.share_institute = new_institute
        # (ردیف legacy بدون هیچ سهم: مثل رفتار قبلی هیچ دلتایی اعمال نمی‌شود و سهم صفر دست‌نخورده
        #  می‌ماند — تا مسیر delete سهمی را به کیف برنگرداند که هرگز از آن کسر نشده بود.)

    elif trans.type == "tuition" or trans.type == "enrollment_payment":
        # (no-op عمدی — رفتار قبلی هم فقط wallet_balance محلی را عوض می‌کرد که در write-back دور ریخته می‌شد.)
        pass

    else:
        # (no-op عمدی — همان دلیل بالا.)
        pass

    # FIX B2 (الگوی H8-P4): اعمال اتمیک دلتاها در SQL — جمع در دیتابیس، نه خواندن-جمع‌زدن در پایتون.
    # COALESCE رفتار قدیمی None→0 را حفظ می‌کند؛ diff می‌تواند منفی هم باشد (کم‌کردن، عین رفتار قبلی).
    if d_teacher or d_institute:
        wallet_rows = (
            db.query(Student)
            .filter(Student.id == trans.student_id)
            .update(
                {
                    Student.wallet_teacher: func.coalesce(Student.wallet_teacher, 0) + d_teacher,
                    Student.wallet_institute: func.coalesce(Student.wallet_institute, 0) + d_institute,
                },
                synchronize_session=False,
            )
        )
        if wallet_rows != 1:
            db.rollback()
            raise HTTPException(status_code=404, detail="دانش‌آموز یافت نشد")

    # ترتیب سینک (مثل بچ ۱/قدم ۱): اول UPDATE خام، بعد refresh برای مقادیر تازه، بعد سینک بالانس کل.
    # (قاعده‌ی واحد: سینک فقط وقتی UPDATE انجام شده؛ وقتی دلتا صفر است total هم دست‌نخورده می‌ماند.)
    if d_teacher or d_institute:
        db.refresh(student)
        # FIX: Bug 18 - derive the total only through the shared wallet helper.
        student.sync_wallet_balance()

    db.commit()
    return {"message": "تراکنش ویرایش و تراز مالی به‌روز شد."}


# ==========================================
# 🔥 API های جدید: مدیریت دانش‌آموز (حذف و تعلیق)
# ==========================================


# 1. حذف کامل دانش‌آموز به همراه سوابق


@router.delete("/admin/students/{id}")
def delete_student(id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    st = db.query(Student).filter(Student.id == id, Student.is_deleted == False).first()
    if not st:
        raise HTTPException(status_code=404, detail="دانش‌آموز یافت نشد")

    # سافت‌دیلیت: ثبت‌نام‌ها، نمرات، حضورها و تراکنش‌ها دست‌نخورده می‌مانند (حفظ سوابق مالی)؛
    # فقط خود شاگرد آرشیو و نشست‌ها + سایه‌های ورود او باطل/حذف می‌شود (آینه‌ی delete_teacher).
    st.is_deleted = True
    shadow_ids = [i for i in (st.user_id, st.parent_user_id) if i]
    if shadow_ids:
        db.query(UserSession).filter(UserSession.user_id.in_(shadow_ids)).delete()
        db.query(User).filter(User.id.in_(shadow_ids)).delete()
        st.user_id = None
        st.parent_user_id = None
    db.commit()
    return {"message": "دانش‌آموز با موفقیت آرشیو شد و سوابق مالی او حفظ گردید."}


# 2. تغییر وضعیت تعلیق (فعال/غیرفعال)


@router.post("/admin/students/{id}/toggle_suspend")
def toggle_suspend_student(id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):  # FIX H10-S: تعلیق فقط ادمین/منشی (الگوی C1)
    st = db.query(Student).filter(Student.id == id, Student.is_deleted == False).first()
    if not st:
        raise HTTPException(status_code=404, detail="دانش‌آموز یافت نشد")

    # تغییر وضعیت (اگر True بود False میشه و برعکس)
    st.is_suspended = not st.is_suspended
    db.commit()

    status_text = "معلق (غیرفعال)" if st.is_suspended else "فعال"
    return {
        "message": f"وضعیت دانش‌آموز به {status_text} تغییر کرد.",
        "is_suspended": st.is_suspended,
    }


# 3. تاریخچه حضور و غیاب دانش‌آموز در یک کلاس خاص
class StudentAttendanceHistoryRequest(BaseModel):
    student_id: int
    course_id: int


@router.post("/test/transaction_logic")
def test_transaction_logic(data: TransactionTestData, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    """تست کامل منطق ایجاد، ویرایش و حذف تراکنش"""
    # FIX (F-C5): این اندپوینت روی دیتای واقعی می‌نویسد (round-trip خود-تمیزشونده) — در
    # پروداکشن غیرفعال، دقیقاً با همان الگوی docs در main.py (M2).
    if os.getenv("ENV", "development").lower() == "production":
        raise HTTPException(status_code=404, detail="یافت نشد")

    # 1. ایجاد تراکنش تستی
    student = db.query(Student).filter(Student.id == data.student_id).first()
    if not student:
        raise HTTPException(404, "دانش‌آموز یافت نشد")

    # ذخیره مقادیر اولیه کیف پول
    initial_wallet_teacher = (
        student.wallet_teacher if student.wallet_teacher is not None else 0
    )
    initial_wallet_institute = (
        student.wallet_institute if student.wallet_institute is not None else 0
    )
    initial_wallet_balance = (
        student.wallet_balance if student.wallet_balance is not None else 0
    )

    # ایجاد تراکنش
    new_trans = Transaction(
        student_id=data.student_id,
        amount=data.amount,
        payment_method=data.payment_method,
        date=data.date,
        description=data.description,
        type="deposit",
        target_wallet=data.target_wallet,
    )
    db.add(new_trans)

    # بروزرسانی کیف پول (مانند submit_payment)
    if data.target_wallet == "teacher":
        student.wallet_teacher = initial_wallet_teacher + data.amount
    elif data.target_wallet == "institute":
        student.wallet_institute = initial_wallet_institute + data.amount
    else:
        # FIX: Bug 18 - derive the total only through the shared wallet helper.
        student.sync_wallet_balance()

    # بروزرسانی کیف پول کل
    final_w_t = student.wallet_teacher if student.wallet_teacher is not None else 0
    final_w_i = student.wallet_institute if student.wallet_institute is not None else 0
    # FIX: Bug 18 - derive the total only through the shared wallet helper.
    student.sync_wallet_balance()

    db.commit()
    db.refresh(new_trans)

    # 2. ویرایش تراکنش (تست منطق update_transaction)
    update_data = TransactionUpdate(
        amount=data.amount + 5000,  # افزایش 5000 تومانی
        description=f"{data.description} (ویرایش شده)",
        date=data.date,
    )

    # شبیه‌سازی منطق update_transaction
    trans_for_update = (
        # FIX: Bug 12 - exclude archived Transaction rows from this active view.
        db.query(Transaction).filter(Transaction.is_deleted == False, Transaction.is_reversed == False).filter(Transaction.id == new_trans.id).first()
    )
    old_amount = trans_for_update.amount
    new_amount = update_data.amount

    trans_for_update.amount = new_amount
    trans_for_update.description = update_data.description
    trans_for_update.date = update_data.date

    diff = new_amount - old_amount

    # بروزرسانی کیف پول
    if trans_for_update.type == "deposit":
        if trans_for_update.target_wallet == "teacher":
            student.wallet_teacher = (student.wallet_teacher or 0) + diff
        elif trans_for_update.target_wallet == "institute":
            student.wallet_institute = (student.wallet_institute or 0) + diff
        else:
            # FIX: Bug 18 - derive the total only through the shared wallet helper.
            student.sync_wallet_balance()

    # بروزرسانی کیف پول کل
    final_w_t = student.wallet_teacher if student.wallet_teacher is not None else 0
    final_w_i = student.wallet_institute if student.wallet_institute is not None else 0
    # FIX: Bug 18 - derive the total only through the shared wallet helper.
    student.sync_wallet_balance()

    db.commit()

    # 3. حذف تراکنش (تست منطق delete_transaction)
    trans_for_delete = (
        # FIX: Bug 12 - exclude archived Transaction rows from this active view.
        db.query(Transaction).filter(Transaction.is_deleted == False, Transaction.is_reversed == False).filter(Transaction.id == new_trans.id).first()
    )

    # بروزرسانی کیف پول قبل از حذف
    if trans_for_delete.type == "deposit":
        if trans_for_delete.target_wallet == "teacher":
            student.wallet_teacher = (
                student.wallet_teacher or 0
            ) - trans_for_delete.amount
        elif trans_for_delete.target_wallet == "institute":
            student.wallet_institute = (
                student.wallet_institute or 0
            ) - trans_for_delete.amount
        else:
            # FIX: Bug 18 - derive the total only through the shared wallet helper.
            student.sync_wallet_balance()

    # بروزرسانی کیف پول کل
    final_w_t = student.wallet_teacher if student.wallet_teacher is not None else 0
    final_w_i = student.wallet_institute if student.wallet_institute is not None else 0
    # FIX: Bug 18 - derive the total only through the shared wallet helper.
    student.sync_wallet_balance()

    db.delete(trans_for_delete)
    db.commit()

    # مقایسه مقادیر نهایی با مقادیر اولیه
    final_wallet_teacher = (
        student.wallet_teacher if student.wallet_teacher is not None else 0
    )
    final_wallet_institute = (
        student.wallet_institute if student.wallet_institute is not None else 0
    )
    final_wallet_balance = (
        student.wallet_balance if student.wallet_balance is not None else 0
    )

    wallet_correct = (
        final_wallet_teacher == initial_wallet_teacher
        and final_wallet_institute == initial_wallet_institute
        and final_wallet_balance == initial_wallet_balance
    )

    return {
        "test_passed": wallet_correct,
        "initial_wallets": {
            "teacher": initial_wallet_teacher,
            "institute": initial_wallet_institute,
            "balance": initial_wallet_balance,
        },
        "final_wallets": {
            "teacher": final_wallet_teacher,
            "institute": final_wallet_institute,
            "balance": final_wallet_balance,
        },
        "test_description": "ایجاد → ویرایش (+5000) → حذف تراکنش",
        "expected_result": "کیف پول باید به مقدار اولیه برگردد",
        "actual_result": "برگشت به مقدار اولیه" if wallet_correct else "عدم تطابق",
    }


@router.post("/sms/send_bulk")
# FIX (L14/R1-align): ارسال گروهی هم ادمین/منشی — هماهنگ با /sms/send تکی (audit-v2/#10) و یادآور قسط (L14/R1). منشی همین قدرت را تک‌تک داشت؛ bulk فقط راحت‌تر و حسابرسی‌پذیرتر است.
def send_bulk_sms(req: BulkSmsRequest, db: Session = Depends(get_db), _: str = Depends(check_admin_or_secretary_access)):
    results = []
    success_count = 0
    today = datetime.datetime.now().strftime("%Y/%m/%d %H:%M")
    
    for s_id in req.student_ids:
        st = db.query(Student).filter(Student.id == s_id, Student.is_deleted == False).first()
        if not st:
            results.append({"student_id": s_id, "status": "failed", "message": "دانش‌آموز یافت نشد"})
            continue
            
        w_t = st.wallet_teacher if st.wallet_teacher is not None else 0
        w_i = st.wallet_institute if st.wallet_institute is not None else 0
        debt = (abs(w_t) if w_t < 0 else 0) + (abs(w_i) if w_i < 0 else 0)
        
        msg = f"سلام {st.first_name} عزیز، بدینوسیله به شما اعلام می‌گردد که مبلغ {debt:,} تومان بدهی شهریه در سیستم دارید. لطفاً جهت تسویه حساب اقدام فرمایید."
        
        # ثبت تک‌تک لاگ‌های پیامک‌ها در دیتابیس
        db.add(
            models.SmsLog(
                target_group=f"student_{st.id}",
                message_text=msg,
                sent_count=1,
                date=today,
            )
        )
        success_count += 1
        results.append({"student_id": s_id, "student_name": f"{st.first_name} {st.last_name}", "status": "success", "message": "پیامک یادآوری بدهی ارسال شد"})
        
    db.commit()
    return {
        "status": "success",
        "message": f"پیامک برای {success_count} دانش‌آموز با موفقیت ارسال شد.",
        "results": results
    }


@router.post("/admin/classes/suspend_bulk")
# DELIBERATE (H10-S): تعلیق گروهی کلاس‌ها عمداً فقط-ادمین می‌ماند — شعاع انفجار (N کلاس به‌جای ۱) و یک‌طرفه‌بودن (فقط True، بدون برگشت) با تکی/گروهیِ پیامک فرق دارد. با bulk-sms هماهنگ نکنید.
def suspend_bulk_classes(req: BulkSuspendRequest, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    results = []
    success_count = 0
    
    for c_id in req.course_ids:
        course = db.query(Course).filter(Course.id == c_id).first()
        if not course:
            results.append({"course_id": c_id, "status": "failed", "message": "کلاس یافت نشد"})
            continue
            
        course.is_suspended = True # معلق کردن کلاس
        success_count += 1
        results.append({"course_id": c_id, "course_title": course.title, "status": "success", "message": "کلاس با موفقیت معلق شد"})
        
    db.commit()
    return {
        "status": "success",
        "message": f"تعداد {success_count} کلاس با موفقیت معلق شدند.",
        "results": results
    }

@router.get("/admin/institute_settings")
def get_institute_settings(db: Session = Depends(get_db), sub_role: str = Depends(check_user_login)):
    # FIX (L14/Y3): تنظیمات — شاگرد/ولی 403؛ معلم نسخه‌ی بدون شماره‌کارت (رسید حضور/فاکتور)؛ کارکنان کامل.
    if sub_role not in ("admin", "secretary", "teacher"):
        raise HTTPException(status_code=403, detail="شما دسترسی لازم برای مشاهده تنظیمات آموزشگاه را ندارید")
    settings = db.query(InstituteSettings).first()
    if not settings:
        settings = InstituteSettings()
        db.add(settings)
        db.commit()
        db.refresh(settings)
    if sub_role == "teacher":
        data = {c.name: getattr(settings, c.name) for c in settings.__table__.columns}
        for _f in ("card_number", "card_number_1", "card_holder_1", "card_number_2", "card_holder_2"):
            data[_f] = None
        return data
    return settings

@router.put("/admin/institute_settings")
def update_institute_settings(data: InstituteSettingsModel, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    settings = db.query(InstituteSettings).first()
    if not settings:
        settings = InstituteSettings()
        db.add(settings)
        
    settings.name = data.name
    settings.address = data.address
    settings.phone = data.phone
    if data.official_email is not None:
        settings.official_email = data.official_email
    if data.footer_text is not None:
        settings.footer_text = data.footer_text
    if data.card_number is not None:
        settings.card_number = data.card_number
    if data.manager_mobile_1 is not None:
        settings.manager_mobile_1 = data.manager_mobile_1
    if data.manager_mobile_2 is not None:
        settings.manager_mobile_2 = data.manager_mobile_2
    if data.card_number_1 is not None:
        settings.card_number_1 = data.card_number_1
    if data.card_holder_1 is not None:
        settings.card_holder_1 = data.card_holder_1
    if data.card_number_2 is not None:
        settings.card_number_2 = data.card_number_2
    if data.card_holder_2 is not None:
        settings.card_holder_2 = data.card_holder_2
    if data.teachers_active is not None:
        settings.teachers_active = data.teachers_active
    if data.live_session_max_minutes is not None:
        settings.live_session_max_minutes = data.live_session_max_minutes
    if data.teacher_settlement_alert_days is not None:
        settings.teacher_settlement_alert_days = data.teacher_settlement_alert_days
        
    db.commit()
    db.refresh(settings)
    return {"message": "تنظیمات آموزشگاه با موفقیت بروزرسانی شد", "settings": settings}

@router.post("/admin/institute_settings/upload_logo")
async def upload_institute_logo(file: UploadFile = File(...), db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    settings = db.query(InstituteSettings).first()
    if not settings:
        settings = InstituteSettings()
        db.add(settings)
        
    # Read file and validate size
    MAX_SIZE = 5 * 1024 * 1024 # 5 MB
    contents = await file.read()
    if len(contents) > MAX_SIZE:
        raise HTTPException(status_code=400, detail="حجم فایل نباید بیشتر از ۵ مگابایت باشد")
        
    # Validate extension
    allowed_extensions = {".jpg", ".jpeg", ".png", ".webp"}
    file_ext = os.path.splitext(file.filename)[1].lower()
    if file_ext not in allowed_extensions:
        raise HTTPException(status_code=400, detail="فرمت فایل غیرمجاز است. فقط فرمت‌های JPG, PNG, WEBP مجاز هستند")
        
    # Generate unique filename
    unique_filename = f"logo_{uuid.uuid4().hex}{file_ext}"
    filepath = os.path.join("uploads/profiles", unique_filename)
    
    # Write to disk
    with open(filepath, "wb") as f_out:
        f_out.write(contents)
        
    # Delete old file
    if settings.logo_path:
        old_path = os.path.join("uploads/profiles", settings.logo_path)
        if os.path.exists(old_path):
            try:
                os.remove(old_path)
            except Exception:
                pass
                
    settings.logo_path = unique_filename
    db.commit()
    
    return {
        "message": "لوگوی آموزشگاه با موفقیت آپلود شد",
        "logo_path": unique_filename,
        "url": f"uploads/profiles/{unique_filename}"
    }

# ==========================================
# آرشیو کلاس‌های حذف‌شده (فقط ادمین) — فقط خواندنی + branch isolation
# ==========================================
def _fmt_datetime(dt) -> str:
    """قالب تاریخ پروژه؛ None → رشته‌ی خالی (هیچ تاریخی حدس زده نمی‌شود)."""
    return dt.strftime("%Y/%m/%d %H:%M") if dt else ""


def _safe_person_name(obj) -> str:
    """نام legacy ممکن است NULL باشد — هرگز «None None» برنگردان."""
    if obj is None:
        return ""
    return f"{obj.first_name or ''} {obj.last_name or ''}".strip()


def _archived_deletion_meta(db: Session, course_ids: List[int]) -> dict:
    """متادیتای حذف از جدول ClassDeletionRequest (بدون تاریخ حدسی).

    `Course` ستون تاریخ حذف **ندارد**؛ هر دو مسیر حذف (حذف مستقیم ادمین در classes.py و
    تایید درخواست) یک ردیف `approved` با `decided_at` ثبت می‌کنند. کلاس‌های legacy که پیش از
    این جدول حذف شده‌اند تاریخ ندارند و مقدارشان خالی می‌ماند.
    """
    if not course_ids:
        return {}
    rows = (
        db.query(models.ClassDeletionRequest)
        .filter(
            models.ClassDeletionRequest.course_id.in_(course_ids),
            models.ClassDeletionRequest.status == "approved",
        )
        .order_by(models.ClassDeletionRequest.id)
        .all()
    )
    meta: dict = {}
    for r in rows:  # آخرین ردیف (بیشترین id) متادیتا را می‌دهد؛ اولین تاریخ موجود حفظ می‌شود.
        ts = r.decided_at or r.created_at
        entry = meta.get(r.course_id)
        if entry is None:
            meta[r.course_id] = {
                "deleted_at": ts,
                "forgive_session_charges": bool(r.forgive_session_charges),
                "requested_by_role": r.requested_by_role,
                "admin_note": r.admin_note,
            }
        else:
            if entry["deleted_at"] is None and ts is not None:
                entry["deleted_at"] = ts
            entry["forgive_session_charges"] = bool(r.forgive_session_charges)
            entry["requested_by_role"] = r.requested_by_role
            entry["admin_note"] = r.admin_note
    return meta


def _archived_class_aggregates(db: Session, course_ids: List[int]) -> dict:
    """شمارش تاریخی ثبت‌نام/جلسه/تراکنش هر کلاس با کوئری گروهی (بدون N+1).

    همه‌ی ردیف‌ها شمرده می‌شوند (آرشیوشده‌ها هم) چون آرشیو باید **تاریخچه** را نشان دهد؛
    `Course.is_deleted`/`Enrollment.is_deleted` هیچ ردیفی را حذف نمی‌کنند.
    """
    stats: dict = {cid: {"students_total": 0, "students_active": 0, "sessions_total": 0,
                         "sessions_archived": 0, "transactions_count": 0, "transactions_total": 0}
                   for cid in course_ids}
    if not course_ids:
        return stats
    for cid, is_deleted, cnt in (
        db.query(Enrollment.course_id, Enrollment.is_deleted, func.count(Enrollment.id))
        .filter(Enrollment.course_id.in_(course_ids))
        .group_by(Enrollment.course_id, Enrollment.is_deleted)
        .all()
    ):
        if cid not in stats:
            continue
        stats[cid]["students_total"] += cnt
        # is_deleted=None برای ردیف‌های legacy یعنی «آرشیو نشده» (سیاست موجود پروژه).
        if not is_deleted:
            stats[cid]["students_active"] += cnt
    for cid, is_deleted, cnt in (
        db.query(SessionLog.course_id, SessionLog.is_deleted, func.count(SessionLog.id))
        .filter(SessionLog.course_id.in_(course_ids))
        .group_by(SessionLog.course_id, SessionLog.is_deleted)
        .all()
    ):
        if cid not in stats:
            continue
        stats[cid]["sessions_total"] += cnt
        if is_deleted:
            stats[cid]["sessions_archived"] += cnt
    for cid, cnt, total in (
        db.query(
            Transaction.course_id,
            func.count(Transaction.id),
            func.coalesce(func.sum(Transaction.amount), 0),
        )
        .filter(Transaction.course_id.in_(course_ids))
        .group_by(Transaction.course_id)
        .all()
    ):
        if cid not in stats:
            continue
        stats[cid]["transactions_count"] = cnt
        stats[cid]["transactions_total"] = total or 0
    return stats


@router.get("/admin/deleted_classes")
def get_deleted_classes(
    query: Optional[str] = None,
    branch_id: Optional[int] = None,
    db: Session = Depends(get_db),
    authorization: Optional[str] = Header(None),
    _: str = Depends(check_admin_access),
):
    """آرشیو کلاس‌های حذف‌شده — فقط ادمین، با branch isolation و اطلاعات کامل.

    تغییرات نسبت به نسخه‌ی قبلی: فیلدهای branch/تعداد دانش‌آموز/جلسه/تراکنش و تاریخ حذف
    اضافه شده، جست‌وجو ممکن است و ترتیب قطعی است. کلیدهای قبلی (id/title/code/teacher_name/
    grade_level/bg_color) دست‌نخورده مانده‌اند تا کلاینت قدیمی نشکند.
    """
    # الگوی branch isolation پروژه (analytics/crm/exports) — ادمینِ شعبه‌دار فقط شعبه‌ی خودش.
    from routers.analytics import get_user_branch_filter  # lazy: جلوگیری از import چرخه‌ای
    resolved_branch = get_user_branch_filter(db, authorization, branch_id)

    # FIX(null-data): is_deleted=True صریح؛ ردیف‌های legacy با NULL «فعال» حساب می‌شوند و
    # طبق سیاست موجود پروژه در آرشیو نمی‌آیند (لیست عادی و آرشیو قاطی نمی‌شوند).
    q = db.query(Course).filter(Course.is_deleted == True)  # noqa: E712
    if resolved_branch is not None:
        q = q.filter(Course.branch_id == resolved_branch)
    if query and query.strip():
        like = f"%{query.strip()}%"
        q = q.outerjoin(Teacher, Teacher.id == Course.teacher_id).filter(
            or_(
                Course.title.ilike(like),
                Course.code.ilike(like),
                Teacher.first_name.ilike(like),
                Teacher.last_name.ilike(like),
            )
        )
    courses = q.all()
    course_ids = [c.id for c in courses]

    meta_map = _archived_deletion_meta(db, course_ids)
    stats_map = _archived_class_aggregates(db, course_ids)

    teacher_ids = {c.teacher_id for c in courses if c.teacher_id is not None}
    branch_ids = {c.branch_id for c in courses if c.branch_id is not None}
    teachers = {t.id: t for t in db.query(Teacher).filter(Teacher.id.in_(teacher_ids)).all()} if teacher_ids else {}
    branches = {b.id: b for b in db.query(models.Branch).filter(models.Branch.id.in_(branch_ids)).all()} if branch_ids else {}

    # ترتیب قطعی: تازه‌ترین حذف اول؛ کلاس‌های بدون تاریخ حذف آخر (شناسه نزولی).
    courses.sort(
        key=lambda c: (meta_map.get(c.id, {}).get("deleted_at") or datetime.datetime.min, c.id),
        reverse=True,
    )

    result = []
    for c in courses:
        st = stats_map.get(c.id, {})
        meta = meta_map.get(c.id, {})
        result.append({
            "id": c.id,
            # FIX(null-data): فیلدهای legacy می‌توانند NULL باشند — رشته‌ی خالی می‌رود و
            # فال‌بک نمایش در اپ انجام می‌شود (نه «None/null» در UI).
            "title": c.title or "",
            "code": c.code or "",
            "teacher_id": c.teacher_id,
            "teacher_name": _safe_person_name(teachers.get(c.teacher_id)) or "نامشخص",
            "branch_id": c.branch_id,
            "branch_name": (branches.get(c.branch_id).name if branches.get(c.branch_id) else "") or "",
            "grade_level": c.grade_level or "",
            "days_of_week": c.days_of_week or "",
            "class_time": c.class_time or "",
            "students_count": st.get("students_total", 0),
            "students_active_count": st.get("students_active", 0),
            "sessions_count": st.get("sessions_total", 0),
            "transactions_count": st.get("transactions_count", 0),
            "deleted_at": _fmt_datetime(meta.get("deleted_at")),
            "forgive_session_charges": bool(meta.get("forgive_session_charges", False)),
            "is_suspended": bool(c.is_suspended),
            "bg_color": c.bg_color or "#FFFFFF",
        })
    return result


@router.get("/admin/deleted_classes/{course_id}")
def get_deleted_class_detail(
    course_id: int,
    db: Session = Depends(get_db),
    authorization: Optional[str] = Header(None),
    _: str = Depends(check_admin_access),
):
    """جزئیات کنترول‌شده‌ی یک کلاس آرشیوشده (فقط ادمین + همان شعبه).

    فقط کلاسِ آرشیوشده سرو می‌شود؛ کلاس فعال، ناموجود و متعلق به شعبه‌ی دیگر همه 404
    می‌گیرند (بدون نشت وجود رکورد). خروجی dict تعریف‌شده است، نه dump خام ORM.
    """
    from routers.analytics import get_user_branch_filter  # lazy: جلوگیری از import چرخه‌ای
    resolved_branch = get_user_branch_filter(db, authorization, None)

    course = (
        db.query(Course)
        .filter(Course.id == course_id, Course.is_deleted == True)  # noqa: E712
        .first()
    )
    if not course or (resolved_branch is not None and course.branch_id != resolved_branch):
        raise HTTPException(status_code=404, detail="کلاس آرشیوشده یافت نشد")

    meta = _archived_deletion_meta(db, [course.id]).get(course.id, {})
    stats = _archived_class_aggregates(db, [course.id]).get(course.id, {})
    teacher = db.query(Teacher).filter(Teacher.id == course.teacher_id).first() if course.teacher_id else None
    branch = db.query(models.Branch).filter(models.Branch.id == course.branch_id).first() if course.branch_id else None
    archived_enrollments = db.query(Enrollment).filter(
        Enrollment.course_id == course.id, Enrollment.is_deleted == True  # noqa: E712
    ).count()
    archived_sessions = db.query(SessionLog).filter(
        SessionLog.course_id == course.id, SessionLog.is_deleted == True  # noqa: E712
    ).count()

    return {
        # مشخصات پایه (همه null-safe)
        "id": course.id,
        "title": course.title or "",
        "code": course.code or "",
        "grade_level": course.grade_level or "",
        "days_of_week": course.days_of_week or "",
        "class_time": course.class_time or "",
        "is_suspended": bool(course.is_suspended),
        "bg_color": course.bg_color or "#FFFFFF",
        # معلم / شعبه
        "teacher_id": course.teacher_id,
        "teacher_name": _safe_person_name(teacher) or "نامشخص",
        "branch_id": course.branch_id,
        "branch_name": (branch.name if branch else "") or "",
        # تاریخچه (حفظ‌شده — چیزی حذف نمی‌شود)
        "students_count": stats.get("students_total", 0),
        "students_active_count": stats.get("students_active", 0),
        "archived_enrollments_count": archived_enrollments,
        "sessions_count": stats.get("sessions_total", 0),
        "archived_sessions_count": archived_sessions,
        "transactions_count": stats.get("transactions_count", 0),
        "transactions_total": stats.get("transactions_total", 0),
        # اطلاعات حذف (از ClassDeletionRequest؛ بدون تاریخ حدسی)
        "deleted_at": _fmt_datetime(meta.get("deleted_at")),
        "has_deletion_record": bool(meta),
        "forgive_session_charges": bool(meta.get("forgive_session_charges", False)),
        "requested_by_role": meta.get("requested_by_role") or "",
        "admin_note": meta.get("admin_note") or "",
    }


@router.get("/admin/pricing_table")
def get_pricing_table(db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    from models import PricingTable
    return db.query(PricingTable).all()


@router.put("/admin/pricing_table")
def update_pricing_table(data: PricingTableUpdateModel, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    from models import PricingTable
    for r in data.rows:
        row = db.query(PricingTable).filter(PricingTable.category == r.category).first()
        if not row:
            row = PricingTable(category=r.category)
            db.add(row)
        row.count_1 = r.count_1
        row.count_2 = r.count_2
        row.count_3 = r.count_3
        row.count_4 = r.count_4
        row.count_5 = r.count_5
    db.commit()
    return {"message": "جدول قیمت‌گذاری با موفقیت بروزرسانی شد"}


@router.get("/admin/students/search", response_model=List[PersonListItem])
def search_admin_students(query: Optional[str] = None, db: Session = Depends(get_db), sub_role: str = Depends(check_user_login)):
    # FIX (L14/Y1): جستجوی سراسری PII شاگردان — فقط کارکنان (ادمین/منشی/معلم). شاگرد/ولی 403.
    if sub_role not in ("admin", "secretary", "teacher"):
        raise HTTPException(status_code=403, detail="شما دسترسی لازم برای جستجوی دانش‌آموزان را ندارید")
    q = db.query(Student).filter(Student.is_deleted == False)
    if query:
        search_fmt = f"%{query}%"
        q = q.filter(
            or_(
                Student.first_name.ilike(search_fmt),
                Student.last_name.ilike(search_fmt),
                Student.national_code.ilike(search_fmt)
            )
        )
    students = q.all()
    return [
        PersonListItem(
            id=s.id,
            # FIX(null-data): هم‌سان با جستجوی معلمان — یک رکورد ناقص نباید کل لیست را 500 کند.
            name=f"{s.first_name or ''} {s.last_name or ''}".strip() or "نامشخص",
            national_code=s.national_code or "",
            mobile=s.student_mobile or "",
            role="student",
            is_suspended=s.is_suspended if s.is_suspended is not None else False
        )
        for s in students
    ]


@router.get("/admin/teachers/search", response_model=List[PersonListItem])
def search_admin_teachers(query: Optional[str] = None, db: Session = Depends(get_db), sub_role: str = Depends(check_user_login)):
    # FIX (L14/Y2): جستجوی سراسری PII معلمان — فقط کارکنان (ادمین/منشی/معلم؛ معلم برای PersonList). شاگرد/ولی 403.
    if sub_role not in ("admin", "secretary", "teacher"):
        raise HTTPException(status_code=403, detail="شما دسترسی لازم برای جستجوی معلمان را ندارید")
    q = db.query(Teacher).filter(Teacher.is_deleted == False)
    if query:
        search_fmt = f"%{query}%"
        q = q.filter(
            or_(
                Teacher.first_name.ilike(search_fmt),
                Teacher.last_name.ilike(search_fmt),
                Teacher.national_code.ilike(search_fmt)
            )
        )
    teachers = q.all()
    return [
        PersonListItem(
            id=t.id,
            # FIX(null-data): فیلدهای legacy ممکن است NULL باشند — پاسخ امن به‌جای ValidationError/500
            # (یک رکورد ناقص نباید کل لیست مربیان/افراد را خراب کند).
            name=f"{t.first_name or ''} {t.last_name or ''}".strip() or "نامشخص",
            national_code=t.national_code or "",
            mobile=t.mobile or "",
            role="teacher",
            is_suspended=t.is_suspended if t.is_suspended is not None else False
        )
        for t in teachers
    ]


import random

# Pydantic models for credentials
class TeacherCredentialsResponse(BaseModel):
    id: int
    name: str
    national_code: str
    mobile: str
    teacher_code: Optional[int] = None
    password: str
    card_number: Optional[str] = None

class TeacherCredentialsUpdateRequest(BaseModel):
    mobile: str
    card_number: Optional[str] = None

@router.get("/admin/teachers/{id}/credentials", response_model=TeacherCredentialsResponse)
def get_teacher_credentials(id: int, db: Session = Depends(get_db), _: str = Depends(check_admin_access)):
    teacher = db.query(Teacher).filter(Teacher.id == id, Teacher.is_deleted == False).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="معلم یافت نشد")
    return TeacherCredentialsResponse(
        id=teacher.id,
        # FIX(null-data): این response_model فیلدهای str اجباری دارد و رکورد legacy با
        # mobile/national_code/نام NULL قبلاً ValidationError/500 می‌داد؛ الان مقدار امن برمی‌گردد.
        name=f"{teacher.first_name or ''} {teacher.last_name or ''}".strip() or "نامشخص",
        national_code=teacher.national_code or "",
        mobile=teacher.mobile or "",
        teacher_code=teacher.teacher_code,
        password=teacher.password or "",
        card_number=teacher.card_number
    )

@router.post("/admin/teachers/{id}/credentials/reset_password")
def reset_teacher_password(id: int, db: Session = Depends(get_db), current_user = Depends(get_current_user), _: str = Depends(check_admin_access)):
    teacher = db.query(Teacher).filter(Teacher.id == id, Teacher.is_deleted == False).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="معلم یافت نشد")
        
    # FIX H12: کد بدون‌تصادم از شمارنده‌ی ترتیبی (نه رندوم خام روی ستون UNIQUE).
    new_code = get_next_sequence_value(db, "teacher", 101)

    # FIX H12: پسورد تصادفی امن با الگوی C12 (مستقل از کد حدس‌زدنی)؛ فقط همین یک‌بار در پاسخ برمی‌گردد.
    _READABLE_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz23456789"  # بدون 0/O/1/l/I
    new_password = "".join(secrets.choice(_READABLE_ALPHABET) for _ in range(8))
    
    old_code = teacher.teacher_code
    old_password = teacher.password
    
    teacher.teacher_code = new_code
    teacher.password = hash_password(new_password)

    # FIX (same shadow-sync as change_password): لاگین اول شاخه‌ی User را چک می‌کند،
    # پس بدون سینک سایه، معلم با پسورد جدید قفل می‌شود (400) — سایه را هم به‌روز کن.
    shadow = db.query(User).filter(User.username == teacher.mobile, User.role == "teacher").first()
    if shadow is None:
        # fallback ناهماهنگی‌های قدیمی (username سایه با موبایل فعلی فرق دارد): حل از طریق سشن
        last_sess = db.query(UserSession).filter(UserSession.teacher_id == teacher.id).order_by(UserSession.id.desc()).first()
        if last_sess is not None:
            cand = db.query(User).filter(User.id == last_sess.user_id, User.role == "teacher").first()
            if cand is not None:
                shadow = cand
    if shadow is not None:
        shadow.password = teacher.password

    # Write to activity log
    log_entry = ActivityLog(
        admin_username=current_user.username,
        action="reset_teacher_password",
        target_id=teacher.id,
        target_name=f"{teacher.first_name or ''} {teacher.last_name or ''}".strip() or "نامشخص",
        details=f"تغییر رمز از {old_code}/{old_password} to {new_code}"
    )
    db.add(log_entry)
    db.commit()
    
    return {
        "message": "رمز عبور جدید معلم با موفقیت صادر شد",
        "new_password": new_password
    }

@router.put("/admin/teachers/{id}/credentials")
def update_teacher_credentials(id: int, data: TeacherCredentialsUpdateRequest, db: Session = Depends(get_db), current_user = Depends(get_current_user), _: str = Depends(check_admin_access)):
    teacher = db.query(Teacher).filter(Teacher.id == id, Teacher.is_deleted == False).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="معلم یافت نشد")

    # FIX H20: نرمال‌سازی به‌جای رد خام H1 — ورودی معتبرِ بدفرمت تمیز و canonical ذخیره می‌شود.
    new_mobile = normalize_mobile(data.mobile)
    if not new_mobile:
        raise HTTPException(status_code=400, detail="فرمت شماره موبایل جدید صحیح نیست. باید ۱۰ یا ۱۱ رقم باشد")

    old_mobile = teacher.mobile
    old_card = teacher.card_number

    # FIX (same pattern as change_mobile): چک یکتایی موبایل جدید بین معلم‌های دیگر، قبل از commit.
    if new_mobile != old_mobile:
        clash = db.query(Teacher).filter(Teacher.mobile == new_mobile, Teacher.id != teacher.id).first()
        if clash:
            raise HTTPException(status_code=409, detail="این شماره موبایل قبلاً توسط معلم دیگری ثبت شده است")

    teacher.mobile = new_mobile
    teacher.card_number = data.card_number

    # FIX: سینک username سایه با موبایل جدید تا لینک سایه یتیم نشود (وگرنه «لاگین زامبی» با موبایل قدیمی).
    if new_mobile != old_mobile:
        shadow = db.query(User).filter(User.username == old_mobile, User.role == "teacher").first()
        if shadow is None:
            # fallback ناهماهنگی‌های قدیمی (username سایه با موبایل فعلی فرق دارد): حل از طریق سشن
            last_sess = db.query(UserSession).filter(UserSession.teacher_id == teacher.id).order_by(UserSession.id.desc()).first()
            if last_sess is not None:
                cand = db.query(User).filter(User.id == last_sess.user_id, User.role == "teacher").first()
                if cand is not None:
                    shadow = cand
        if shadow is not None:
            shadow.username = new_mobile

    # Write to activity log
    log_entry = ActivityLog(
        admin_username=current_user.username,
        action="update_teacher_credentials",
        target_id=teacher.id,
        target_name=f"{teacher.first_name or ''} {teacher.last_name or ''}".strip() or "نامشخص",
        details=f"موبایل: {old_mobile}->{new_mobile} | شماره کارت: {old_card}->{data.card_number}"
    )
    db.add(log_entry)
    from sqlalchemy.exc import IntegrityError  # local: فقط همین‌جا لازم است
    try:
        db.commit()
    except IntegrityError:
        # backstop مسابقه‌ی هم‌زمان (race) بین چک بالا و commit
        db.rollback()
        raise HTTPException(status_code=409, detail="این شماره موبایل قبلاً در سیستم ثبت شده است")
    
    return {
        "message": "اطلاعات تماس و حساب بانکی معلم با موفقیت به‌روزرسانی شد"
    }
