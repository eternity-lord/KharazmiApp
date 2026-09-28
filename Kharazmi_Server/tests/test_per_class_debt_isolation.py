"""Regression coverage for per-enrollment financial isolation.

A student's wallets are aggregate legacy balances. They must not be copied into
an unrelated class row when the student is enrolled in more than one class.
"""
import asyncio
import datetime
import io
import unittest

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from models import (
    Base,
    Branch,
    Course,
    Enrollment,
    SessionLog,
    Student,
    Teacher,
    Transaction,
    User,
    UserSession,
)
from financial_calculations import calculate_enrollment_debt_breakdown, calculate_student_debt
from routers.classes import (
    get_all_classes,
    get_class_details,
    get_class_full_report,
    get_class_students_excel,
    get_class_students_full,
)
from routers.admin import get_student_full_profile
from routers.finance import (
    get_invoice_details,
    get_student_class_status,
    get_student_financial_dashboard,
    search_finance_advanced,
)
from routers.teachers import get_my_classes


class TestPerClassDebtIsolation(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(bind=self.engine)
        self.db = sessionmaker(bind=self.engine, expire_on_commit=False)()

        self.db.add(Branch(id=1, name="مرکزی", active=True))
        self.admin = User(
            id=1, username="admin", password="x", full_name="مدیر",
            role="admin", sub_role="admin", branch_id=1,
        )
        self.teacher = Teacher(
            id=1, first_name="معلم", last_name="مشترک", mobile="09120000001",
            national_code="0012345678", is_approved=True,
        )
        self.student = Student(
            id=1, first_name="علی", last_name="دوکلاسه", national_code="0012345680",
            student_mobile="09121111111", wallet_teacher=-300,
            wallet_institute=-200, wallet_balance=-500,
        )
        self.course_a = Course(
            id=1, title="کلاس A", code="A", teacher_id=self.teacher.id,
            branch_id=1, teacher_session_price=100, is_admin_approved=True,
            is_deleted=False,
        )
        self.course_b = Course(
            id=2, title="کلاس B", code="B", teacher_id=self.teacher.id,
            branch_id=1, teacher_session_price=100, is_admin_approved=True,
            is_deleted=False,
        )
        self.enrollment_a = Enrollment(
            id=1, student_id=self.student.id, course_id=self.course_a.id,
            branch_id=1, register_date="1405/06/01", shift="عصر",
            total_tuition=1000, total_paid=0, is_deleted=False,
        )
        # B intentionally has no priced tuition: this is the legacy enrollment
        # that used to inherit A's aggregate wallet debt.
        self.enrollment_b = Enrollment(
            id=2, student_id=self.student.id, course_id=self.course_b.id,
            branch_id=1, register_date="1405/06/02", shift="عصر",
            total_tuition=0, total_paid=0, is_deleted=False,
        )
        self.db.add_all([
            self.admin, self.teacher, self.student, self.course_a, self.course_b,
            self.enrollment_a, self.enrollment_b,
            UserSession(
                token="admin-token", user_id=self.admin.id, sub_role="admin",
                created_at=datetime.datetime.now(),
            ),
            # A session charge is class-scoped and creates the legacy wallet
            # state represented above; it is never a payment for class B.
            SessionLog(id=1, course_id=self.course_a.id, date="1405/06/03", session_code=1),
            Transaction(
                id=1, student_id=self.student.id, course_id=self.course_a.id,
                session_id=1, amount=-500, type="session_charge",
                share_teacher=300, share_institute=200,
                date="1405/06/03", description="A charge",
                is_deleted=False, is_reversed=False,
            ),
        ])
        self.db.commit()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(bind=self.engine)
        self.engine.dispose()

    def _admin(self):
        return "Bearer admin-token"

    def _excel_rows(self, course_id):
        response = get_class_students_excel(class_id=course_id, db=self.db, _="admin")

        async def read_body():
            return b"".join([chunk async for chunk in response.body_iterator])

        payload = asyncio.run(read_body())
        sheet = load_workbook(io.BytesIO(payload), data_only=True).active
        return list(sheet.iter_rows(min_row=4, values_only=True))

    def test_breakdown_does_not_copy_wallet_debt_to_legacy_class(self):
        a = calculate_enrollment_debt_breakdown(self.db, self.enrollment_a)
        b = calculate_enrollment_debt_breakdown(self.db, self.enrollment_b)
        self.assertEqual(a["debt"], 1000)
        self.assertEqual(a["debt_teacher"], 300)
        self.assertEqual(a["debt_institute"], 200)
        self.assertEqual(b["debt"], 0)
        self.assertEqual(b["debt_teacher"], 0)
        self.assertEqual(b["debt_institute"], 0)

    def test_class_list_details_full_report_and_excel_are_isolated(self):
        classes = {row["id"]: row for row in get_all_classes(db=self.db, _="admin")}
        self.assertEqual(classes[self.course_a.id]["total_debt"], 1000)
        self.assertEqual(classes[self.course_b.id]["total_debt"], 0)
        self.assertEqual(classes[self.course_b.id]["debt_to_teacher"], 0)
        self.assertEqual(classes[self.course_b.id]["debt_to_institute"], 0)

        details_b = get_class_details(
            course_id=self.course_b.id, db=self.db,
            authorization=self._admin(), sub_role="admin",
        )["students"][0]
        self.assertEqual(details_b["debt"], 0)
        self.assertEqual(details_b["paid"], 0)

        report_b = get_class_full_report(
            id=self.course_b.id, db=self.db,
            authorization=self._admin(), sub_role="admin",
        )
        self.assertEqual(report_b.info.total_debt, 0)
        self.assertEqual(report_b.info.total_revenue, 0)

        full_b = get_class_students_full(
            id=self.course_b.id, db=self.db,
            authorization=self._admin(), sub_role="admin",
        )["students"][0]
        self.assertEqual(full_b["debt"], 0)
        self.assertEqual(full_b["debt_teacher"], 0)
        self.assertEqual(full_b["debt_institute"], 0)
        # These keys remain aggregate wallets for backward compatibility, not
        # class debt inputs.
        self.assertEqual(full_b["wallet_teacher"], -300)
        self.assertEqual(full_b["wallet_institute"], -200)

        excel_b = self._excel_rows(self.course_b.id)[0]
        # columns: id, name, base, discount type, selected discount,
        # discount amount, final tuition, paid, teacher debt, institute debt,
        # total debt, ...
        self.assertEqual(excel_b[7], 0)
        self.assertEqual(excel_b[8], 0)
        self.assertEqual(excel_b[9], 0)
        self.assertEqual(excel_b[10], 0)

        # finance.py invoice/status/search paths must use the same enrollment scope.
        status_b = get_student_class_status(
            student_id=self.student.id, course_id=self.course_b.id,
            db=self.db, authorization=self._admin(), sub_role="admin",
        )
        self.assertEqual(status_b["due_to_teacher"], 0)
        self.assertEqual(status_b["due_to_institute"], 0)
        invoice_b = get_invoice_details(
            enrollment_id=self.enrollment_b.id, db=self.db,
            authorization=self._admin(), _role="admin",
        )
        self.assertEqual(invoice_b["balance_due"], 0)
        dashboard = get_student_financial_dashboard(
            student_id=self.student.id, db=self.db,
            authorization=self._admin(), _role="admin",
        )
        dashboard_rows = {row["enrollment_id"]: row for row in dashboard["enrollments"]}
        self.assertEqual(dashboard_rows[self.enrollment_b.id]["outstanding"], 0)
        advanced = search_finance_advanced(
            query="کلاس B", branch_id=None, authorization=self._admin(),
            db=self.db, sub_role="admin",
        )
        class_result = next(row for row in advanced if row.type == "class")
        self.assertEqual(class_result.students_in_class[0]["total_debt"], 0)

        teacher_classes = get_my_classes(
            teacher_id=self.teacher.id, db=self.db,
            authorization=self._admin(), sub_role="admin",
        )
        teacher_rows = {row["id"]: row for row in teacher_classes}
        self.assertEqual(teacher_rows[self.course_a.id]["total_debt"], 1000)
        self.assertEqual(teacher_rows[self.course_b.id]["total_debt"], 0)
        self.assertEqual(teacher_rows[self.course_b.id]["debt_to_teacher"], 0)
        self.assertEqual(teacher_rows[self.course_b.id]["debt_to_institute"], 0)

    def test_payment_for_b_does_not_change_a_and_legacy_course_payment_is_scoped(self):
        before_a = calculate_enrollment_debt_breakdown(self.db, self.enrollment_a)
        self.db.add(Transaction(
            id=2, student_id=self.student.id, course_id=self.course_b.id,
            enrollment_id=None, amount=100, target_wallet="institute",
            type="deposit", date="1405/06/04", description="legacy B payment",
            is_deleted=False, is_reversed=False,
        ))
        self.db.commit()
        after_a = calculate_enrollment_debt_breakdown(self.db, self.enrollment_a)
        after_b = calculate_enrollment_debt_breakdown(self.db, self.enrollment_b)
        self.assertEqual(after_a, before_a)
        self.assertEqual(after_b["paid_institute"], 100)
        self.assertEqual(after_b["debt"], 0)

        self.db.add(Transaction(
            id=3, student_id=self.student.id, course_id=None,
            enrollment_id=None, amount=700, target_wallet="institute",
            type="deposit", date="1405/06/04", description="general credit",
            is_deleted=False, is_reversed=False,
        ))
        self.db.commit()
        self.assertEqual(calculate_enrollment_debt_breakdown(self.db, self.enrollment_a), after_a)
        self.assertEqual(calculate_enrollment_debt_breakdown(self.db, self.enrollment_b), after_b)

        classes = {row["id"]: row for row in get_all_classes(db=self.db, _="admin")}
        self.assertEqual(classes[self.course_a.id]["total_debt"], 1000)
        self.assertEqual(classes[self.course_a.id]["total_paid"], 0)
        self.assertEqual(classes[self.course_b.id]["total_paid"], 100)
        self.assertEqual(get_class_details(
            course_id=self.course_a.id, db=self.db,
            authorization=self._admin(), sub_role="admin",
        )["students"][0]["paid"], 0)
        self.assertEqual(get_class_details(
            course_id=self.course_b.id, db=self.db,
            authorization=self._admin(), sub_role="admin",
        )["students"][0]["paid"], 100)
        self.assertEqual(get_class_full_report(
            id=self.course_a.id, db=self.db,
            authorization=self._admin(), sub_role="admin",
        ).info.total_revenue, 0)
        self.assertEqual(get_class_full_report(
            id=self.course_b.id, db=self.db,
            authorization=self._admin(), sub_role="admin",
        ).info.total_revenue, 100)
        self.assertEqual(self._excel_rows(self.course_a.id)[0][7], 0)
        self.assertEqual(self._excel_rows(self.course_b.id)[0][7], 100)

    def test_both_payment_is_split_in_admin_teacher_financial_profile(self):
        self.db.add(Transaction(
            id=4, student_id=self.student.id, course_id=self.course_b.id,
            enrollment_id=self.enrollment_b.id, amount=100, target_wallet="both",
            share_teacher=40, share_institute=60, type="deposit",
            date="1405/06/05", description="both B payment",
            is_deleted=False, is_reversed=False,
        ))
        self.db.commit()
        profile = get_student_full_profile(
            id=self.student.id, authorization=self._admin(), db=self.db, role="admin",
        )
        financial = {row["course_title"]: row for row in profile["teachers_financial"]}
        self.assertEqual(financial["کلاس A"]["paid_teacher"], 0)
        self.assertEqual(financial["کلاس A"]["paid_institute"], 0)
        self.assertEqual(financial["کلاس B"]["paid_teacher"], 40)
        self.assertEqual(financial["کلاس B"]["paid_institute"], 60)
        self.assertEqual(profile["total_paid_institute"], 60)

    def test_class_debt_plus_unallocated_legacy_debt_equals_student_total(self):
        total = calculate_student_debt(self.db, self.student)
        class_sum = sum(
            calculate_enrollment_debt_breakdown(self.db, enrollment)["debt"]
            for enrollment in (self.enrollment_a, self.enrollment_b)
        )
        unallocated = total - class_sum
        self.assertEqual(total, 1000)
        self.assertEqual(class_sum + unallocated, total)
        self.assertGreaterEqual(unallocated, 0)


if __name__ == "__main__":
    unittest.main()
